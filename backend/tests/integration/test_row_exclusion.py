"""⑤ #13 行の除外（② FUNC-05）。

除外した行は未確認から外れ、一括確認の N からも外れ、Excel にも出ない。
**確認済みかどうかは触らない**ので、取り消すと除外する前の状態に戻る。
"""
import uuid

import pytest
from sqlalchemy import delete, select

from app.core.database import AsyncSessionLocal
from app.models import Inquiry, InquiryInput, ItemRow, ItemValue
from app.services import confirm_service
from app.services.dev_fixtures import register_d1
from app.services.item_draft_service import replace_rows
from app.services.item_list_service import (
    _sampling_counts,
    check_row,
    exclude_row,
    get_items,
)
from app.services.source_excerpt_service import get_source
from tests.integration.test_check_completion import _good_row


@pytest.fixture
async def three_rows():
    """確信が高い行を3行持つ案件。"""
    async with AsyncSessionLocal() as session:
        inquiry = await register_d1(session)
        await session.commit()
        excel = (
            await session.execute(
                select(InquiryInput).where(
                    InquiryInput.inquiry_id == inquiry.id, InquiryInput.format == "excel"
                )
            )
        ).scalar_one()
        await replace_rows(session, inquiry.id, [_good_row(excel.id, n) for n in range(1, 4)])
        inquiry.status = "awaiting_review"
        await session.commit()
        rows = (
            (
                await session.execute(
                    select(ItemRow).where(ItemRow.inquiry_id == inquiry.id).order_by(ItemRow.row_no)
                )
            )
            .scalars()
            .all()
        )
        yield inquiry.id, [r.id for r in rows]
    async with AsyncSessionLocal() as session:
        await session.execute(delete(Inquiry).where(Inquiry.id == inquiry.id))
        await session.commit()


async def _open_source(session, inquiry_id: uuid.UUID, row_id: uuid.UUID) -> None:
    value = (
        await session.execute(
            select(ItemValue).where(ItemValue.item_row_id == row_id, ItemValue.field == "item_name")
        )
    ).scalar_one()
    await get_source(session, inquiry_id, value.id)


async def test_excluded_row_drops_out_of_unchecked(three_rows) -> None:
    inquiry_id, row_ids = three_rows
    async with AsyncSessionLocal() as session:
        before = await get_items(session, inquiry_id)
        assert (before["summary"]["unchecked_rows"], before["summary"]["excluded_rows"]) == (3, 0)

        result = await exclude_row(session, inquiry_id, row_ids[0], True)
        await session.commit()
    # 人が見て判断したので、未確認から外れる
    assert (result["summary"]["unchecked_rows"], result["summary"]["excluded_rows"]) == (2, 1)


async def test_excluded_row_drops_out_of_sampling_n(three_rows) -> None:
    """一括確認の N は「確信が高く、除外していない行」。除外すると分母が減る。"""
    inquiry_id, row_ids = three_rows
    async with AsyncSessionLocal() as session:
        assert (await _sampling_counts(session, inquiry_id))["required_samples"] == 3
        await exclude_row(session, inquiry_id, row_ids[0], True)
        await session.commit()
    async with AsyncSessionLocal() as session:
        counts = await _sampling_counts(session, inquiry_id)
    assert (counts["confident_rows"], counts["required_samples"]) == (2, 2)


async def test_cancel_restores_the_previous_check_state(three_rows) -> None:
    """確認済みの行を除外して取り消すと、**確認済みのまま**戻る。"""
    inquiry_id, row_ids = three_rows
    async with AsyncSessionLocal() as session:
        await _open_source(session, inquiry_id, row_ids[0])
        await check_row(session, inquiry_id, row_ids[0])
        await exclude_row(session, inquiry_id, row_ids[0], True)
        await session.commit()
        row = await session.get(ItemRow, row_ids[0])
        assert (row.check_state, row.excluded_at is not None) == ("checked", True)

        result = await exclude_row(session, inquiry_id, row_ids[0], False)
        await session.commit()
    assert result["check_state"] == "checked"
    assert result["summary"]["excluded_rows"] == 0
    # 未確認だった行を除外して戻すと、未確認のまま
    async with AsyncSessionLocal() as session:
        await exclude_row(session, inquiry_id, row_ids[1], True)
        await session.commit()
        back = await exclude_row(session, inquiry_id, row_ids[1], False)
        await session.commit()
    assert back["check_state"] == "unchecked"


async def test_excluded_row_is_not_exported(three_rows) -> None:
    """除外した行は Excel に出ず、item_list_exports.row_count にも含めない。"""
    inquiry_id, row_ids = three_rows
    async with AsyncSessionLocal() as session:
        await exclude_row(session, inquiry_id, row_ids[2], True)
        for row_id in row_ids[:2]:
            await _open_source(session, inquiry_id, row_id)
            await check_row(session, inquiry_id, row_id)
        await session.commit()

    async with AsyncSessionLocal() as session:
        result = await confirm_service.confirm(session, inquiry_id)
        await session.commit()
    assert result["row_count"] == 2

    from openpyxl import load_workbook

    async with AsyncSessionLocal() as session:
        export = await confirm_service.get_export(session, inquiry_id)
    sheet = load_workbook(export.storage_path).active
    body = [r for r in sheet.iter_rows(min_row=4, values_only=True)]
    assert len(body) == 2


async def test_needs_confirmation_filter_skips_excluded(three_rows) -> None:
    """SCR-08 は除外した行を返さない（出力されないので問い合わせの対象にしない）。"""
    inquiry_id, row_ids = three_rows
    async with AsyncSessionLocal() as session:
        # 1行を要確認にしてから除外する
        value = (
            await session.execute(
                select(ItemValue).where(
                    ItemValue.item_row_id == row_ids[0], ItemValue.field == "note"
                )
            )
        ).scalar_one_or_none()
        row = await session.get(ItemRow, row_ids[0])
        row.classification = "needs_confirmation"
        if value is not None:
            value.state = "needs_confirmation"
            value.raw_text = None
            value.source_input_id = None
            value.locator = None
        await exclude_row(session, inquiry_id, row_ids[0], True)
        await session.commit()

    async with AsyncSessionLocal() as session:
        filtered = await get_items(session, inquiry_id, filter_="needs_confirmation")
        with_excluded = await get_items(session, inquiry_id)
    assert [r["row_id"] for r in filtered["rows"]] == []
    # 既定では除外した行も返す（SCR-05 が最後にまとめて表示するため）
    assert str(row_ids[0]) in [r["row_id"] for r in with_excluded["rows"]]
