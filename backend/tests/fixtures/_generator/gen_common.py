# -*- coding: utf-8 -*-
"""引合書ダミー生成の共通部品。"""
import os, random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Pt, Cm

FONT = "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf"
pdfmetrics.registerFont(TTFont("IPAGothic", FONT))

THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDE7F0")


def xlsx_table(path, title_lines, headers, rows, header_row=None, sheet_name="明細",
               col_widths=None, footer_lines=None):
    """よくある見積依頼 Excel。title_lines のあと空行、見出し、明細。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    r = 1
    for line in title_lines:
        ws.cell(row=r, column=1, value=line).font = Font(bold=(r == 1), size=14 if r == 1 else 11)
        r += 1
    r += 1
    if header_row:                      # 見出し行を指定の行まで下げる
        r = header_row
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True)
        cell.border = BOX
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
    hr = r
    for row in rows:
        r += 1
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
    for i, w in enumerate(col_widths or [6, 30, 20, 10, 8, 16, 24], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    if footer_lines:
        r += 2
        for line in footer_lines:
            ws.cell(row=r, column=1, value=line)
            r += 1
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return hr


def pdf_doc(path, title, meta_lines, headers=None, rows=None, notes=None, col_x=None):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    c = rl_canvas.Canvas(path, pagesize=A4)
    W, H = A4
    y = H - 60
    c.setFont("IPAGothic", 16)
    c.drawString(50, y, title); y -= 30
    c.setFont("IPAGothic", 10)
    for line in meta_lines:
        c.drawString(50, y, line); y -= 16
    y -= 10
    if headers:
        xs = col_x or [50, 90, 250, 350, 400, 450]
        c.setFont("IPAGothic", 10)
        for x, h in zip(xs, headers):
            c.drawString(x, y, h)
        y -= 4
        c.line(45, y, W - 45, y); y -= 16
        for row in rows:
            for x, v in zip(xs, row):
                c.drawString(x, y, str(v))
            y -= 16
            if y < 90:
                c.showPage(); y = H - 60; c.setFont("IPAGothic", 10)
    if notes:
        y -= 14
        c.setFont("IPAGothic", 10)
        for line in notes:
            c.drawString(50, y, line); y -= 15
            if y < 60:
                c.showPage(); y = H - 60; c.setFont("IPAGothic", 10)
    c.save()


def docx_table(path, title, meta_lines, headers, rows, notes=None):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "IPAGothic"; st.font.size = Pt(10.5)
    d.add_heading(title, level=1)
    for line in meta_lines:
        d.add_paragraph(line)
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for line in (notes or []):
        d.add_paragraph(line)
    d.save(path)


def docx_bullets(path, title, meta_lines, blocks, notes=None):
    """表ではなく箇条書きで明細を書く『初見の書式』。"""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "IPAGothic"; st.font.size = Pt(10.5)
    d.add_heading(title, level=1)
    for line in meta_lines:
        d.add_paragraph(line)
    d.add_paragraph("")
    for i, b in enumerate(blocks, start=1):
        p = d.add_paragraph(); p.add_run(f"【{i}】{b['name']}").bold = True
        for k in ("型番", "数量", "単位", "希望納期", "備考"):
            if b.get(k):
                d.add_paragraph(f"　{k}：{b[k]}")
        d.add_paragraph("")
    for line in (notes or []):
        d.add_paragraph(line)
    d.save(path)


def write_text(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
