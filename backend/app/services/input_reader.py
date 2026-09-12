"""入力（ファイル・メール本文）を「位置つき」で読む。

読み取り元の位置の書き方は ⑤ の locator_label と同じにそろえる:
| 形式 | 位置の形 | 例 |
|------|---------|-----|
| Excel | シート名!セル番地 | 明細!C12 |
| PDF | p.ページ/行番号 | p.2/14 |
| Word | 表N/行/列 または 段落N | 表1/5/3・段落12 |
| メール本文 | 行番号 | L8 |

列の意味づけ（どの列が型番かなど）は LLM の仕事なので、ここでは整形しない
（agent.md「エージェントに任せること／通常のコードで作ること」）。
"""
from pathlib import Path
from typing import Any

from app.models import InquiryInput

# 1回に返す量の上限（コンテキストを守る）
MAX_CELLS = 300
DEFAULT_LIMIT = 60


class InputNotReadableError(RuntimeError):
    """開けない・壊れている入力（② FUNC-01 の破損・FUNC-07 の判読不能）。"""


def _result(
    cells: list[dict[str, str]],
    *,
    has_more: bool,
    has_readable_text: bool,
    next_: dict[str, Any] | None,
    summary: dict[str, Any],
) -> dict[str, Any]:
    return {
        "cells": cells[:MAX_CELLS],
        "has_more": has_more,
        "has_readable_text": has_readable_text,
        "next": next_,
        "summary": summary,
    }


# ---------------- Excel ----------------
def _read_excel(path: Path, *, sheet: str | None, start: int, limit: int) -> dict[str, Any]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # 壊れたファイル
        raise InputNotReadableError(f"Excel を開けない: {e}") from e

    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    sheet_name = ws.title
    cells: list[dict[str, str]] = []
    last_row = start - 1
    rows_read = 0

    for offset, row in enumerate(ws.iter_rows(min_row=start, max_row=start + limit - 1)):
        rows_read += 1
        # read_only では空セルが EmptyCell（row 属性なし）になるため、行番号は数え上げで持つ
        last_row = start + offset
        for cell in row:
            if cell.value is None or str(cell.value).strip() == "":
                continue
            cells.append({"locator": f"{sheet_name}!{cell.coordinate}", "text": str(cell.value)})

    max_row = ws.max_row or last_row
    has_more = last_row < max_row and rows_read == limit
    wb.close()
    return _result(
        cells,
        has_more=has_more,
        has_readable_text=bool(cells),
        next_={"sheet": sheet_name, "start": last_row + 1} if has_more else None,
        summary={
            "file": path.name,
            "range": f"{sheet_name}!{start}〜{last_row}行",
            "cells": len(cells),
            "has_more": has_more,
        },
    )


# ---------------- メール本文 ----------------
def _read_mail_body(text: str, *, display_name: str, start: int, limit: int) -> dict[str, Any]:
    lines = text.splitlines()
    chunk = lines[start - 1 : start - 1 + limit]
    cells = [
        {"locator": f"L{start + i}", "text": line} for i, line in enumerate(chunk) if line.strip()
    ]
    end = start - 1 + len(chunk)
    has_more = end < len(lines)
    return _result(
        cells,
        has_more=has_more,
        has_readable_text=bool(cells),
        next_={"start": end + 1} if has_more else None,
        summary={
            "file": display_name,
            "range": f"L{start}〜L{end}",
            "lines": len(cells),
            "has_more": has_more,
        },
    )


# ---------------- Word ----------------
def _read_word(path: Path, *, start: int, limit: int) -> dict[str, Any]:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as e:
        raise InputNotReadableError(f"Word を開けない: {e}") from e

    items: list[dict[str, str]] = []
    for idx, para in enumerate(document.paragraphs, start=1):
        if para.text.strip():
            items.append({"locator": f"段落{idx}", "text": para.text})
    for t_idx, table in enumerate(document.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            for c_idx, cell in enumerate(row.cells, start=1):
                if cell.text.strip():
                    items.append({"locator": f"表{t_idx}/{r_idx}/{c_idx}", "text": cell.text})

    chunk = items[start - 1 : start - 1 + limit]
    end = start - 1 + len(chunk)
    has_more = end < len(items)
    return _result(
        chunk,
        has_more=has_more,
        has_readable_text=bool(chunk),
        next_={"start": end + 1} if has_more else None,
        summary={
            "file": path.name,
            "range": f"{start}〜{end} 番目の段落/セル",
            "items": len(chunk),
            "has_more": has_more,
        },
    )


# ---------------- PDF ----------------
def _read_pdf(path: Path, *, page: int, start: int, limit: int) -> dict[str, Any]:
    import pdfplumber

    try:
        pdf = pdfplumber.open(path)
    except Exception as e:
        raise InputNotReadableError(f"PDF を開けない: {e}") from e

    with pdf:
        total_pages = len(pdf.pages)
        if page > total_pages:
            return _result(
                [],
                has_more=False,
                has_readable_text=False,
                next_=None,
                summary={"file": path.name, "range": f"p.{page}（ページなし）", "lines": 0},
            )
        text = pdf.pages[page - 1].extract_text() or ""

    lines = [ln for ln in text.splitlines()]
    chunk = lines[start - 1 : start - 1 + limit]
    cells = [
        {"locator": f"p.{page}/{start + i}", "text": line}
        for i, line in enumerate(chunk)
        if line.strip()
    ]
    end = start - 1 + len(chunk)
    more_in_page = end < len(lines)
    has_more = more_in_page or page < total_pages
    if more_in_page:
        next_ = {"page": page, "start": end + 1}
    elif page < total_pages:
        next_ = {"page": page + 1, "start": 1}
    else:
        next_ = None

    return _result(
        cells,
        has_more=has_more,
        # 文字が取れない = 判読不能の手がかり（② FUNC-07）
        has_readable_text=bool(cells),
        next_=next_,
        summary={
            "file": path.name,
            "range": f"p.{page}/{start}〜{end}（全{total_pages}ページ）",
            "lines": len(cells),
            "has_more": has_more,
        },
    )


def read_input(
    input_: InquiryInput,
    *,
    sheet: str | None = None,
    page: int = 1,
    start: int = 1,
    limit: int = DEFAULT_LIMIT,
) -> dict[str, Any]:
    """入力を位置つきで読む。形式ごとに位置の書き方が変わる。"""
    limit = max(1, min(limit, MAX_CELLS))
    start = max(1, start)

    if input_.kind == "mail_body":
        return _read_mail_body(
            input_.content_text or "", display_name=input_.display_name, start=start, limit=limit
        )

    path = Path(input_.storage_path or "")
    if not path.exists():
        raise InputNotReadableError(f"原本が見つからない: {path}")

    if input_.format == "excel":
        return _read_excel(path, sheet=sheet, start=start, limit=limit)
    if input_.format == "word":
        return _read_word(path, start=start, limit=limit)
    if input_.format == "pdf":
        return _read_pdf(path, page=max(1, page), start=start, limit=limit)
    raise InputNotReadableError(f"対応しない形式: {input_.format}")
