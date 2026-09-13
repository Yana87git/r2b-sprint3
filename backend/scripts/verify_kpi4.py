"""KPI4 の「見逃し0件」を検証する: 誤りを直して確定し、出力 Excel に誤りが残らないか見る。

  uv run python scripts/verify_kpi4.py <case>=<inquiry_id> ...
  uv run python scripts/verify_kpi4.py c2=8dc8c412-... c3=...

**画面と同じ API（#9 → #10 → #11 → #17 → #18）だけを使う。** 人がやる操作をそのまま辿る。
時間（1件10分以内）は人が測る。ここで見るのは「誤りが出力に残っていないか」だけ。

C7・C8（H1 入力間の不一致）の扱い: どちらが正しいかは書類から決まらないので、
**両方の行を残したまま、食い違っている項目を「要確認（顧客回答待ち）」にする**。
確かでない数字を確定した値として外に出さない、という判断（① ビジネスゴール2）。
"""
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import httpx
from openpyxl import load_workbook

BASE = "http://localhost:8000"
PENDING = "顧客回答待ち"


def norm(text: Any) -> str:
    return unicodedata.normalize("NFKC", str(text)).replace(" ", "").lower()


@dataclass
class Case:
    key: str
    note: str
    # 品目リスト案を受け取り、{row_id: {field: 更新内容}} を返す
    fix: Callable[[list[dict]], dict[str, dict]]
    # 出力 Excel の全セルを受け取り、(合格か, 説明) を返す
    judge: Callable[[list[tuple]], tuple[bool, str]]


def _row_of(rows: list[dict], text: str) -> dict:
    """値のどれかに text を含む行（原文でも正規化後でもよい）。"""
    needle = norm(text)
    for row in rows:
        for value in row["values"].values():
            if needle in norm(value.get("raw_text") or "") or needle in norm(
                value.get("value_text") or ""
            ):
                return row
    raise LookupError(text)


def _rows_of(rows: list[dict], text: str) -> list[dict]:
    needle = norm(text)
    found = []
    for row in rows:
        if any(
            needle in norm(v.get("raw_text") or "") or needle in norm(v.get("value_text") or "")
            for v in row["values"].values()
        ):
            found.append(row)
    return found


def _cells_text(cells: list[tuple]) -> str:
    return " ".join(norm(c) for row in cells for c in row if c is not None)


# 出力 Excel の列: No / 品目名 / 型番 / 数量 / 単位 / 納期 / 備考 / 状態
UNIT_COLUMN = 4


def _unit_is_not(wrong_unit: str) -> Callable[[list[tuple]], tuple[bool, str]]:
    """単位の誤り（m・kg）が1行も残っていないか。"""

    def judge(cells: list[tuple]) -> tuple[bool, str]:
        leftovers = [
            row
            for row in cells
            if len(row) > UNIT_COLUMN and norm(row[UNIT_COLUMN]) == norm(wrong_unit)
        ]
        return (
            not leftovers,
            f"単位が「{wrong_unit}」の行が{len(leftovers)}行"
            if leftovers
            else f"単位「{wrong_unit}」の行は無い",
        )

    return judge


def _absent(marker: str) -> Callable[[list[tuple]], tuple[bool, str]]:
    def judge(cells: list[tuple]) -> tuple[bool, str]:
        present = norm(marker) in _cells_text(cells)
        return (not present, f"出力に「{marker}」が{'残っている' if present else '無い'}")

    return judge


def _fix_value(rows: list[dict], marker: str, field: str, payload: dict) -> dict[str, dict]:
    return {_row_of(rows, marker)["row_id"]: {field: payload}}


def _flag_all(rows: list[dict], marker: str, field: str) -> dict[str, dict]:
    """食い違っている項目を、両方の行とも「要確認」にする（C7・C8）。"""
    return {
        row["row_id"]: {field: {"state": "needs_confirmation"}} for row in _rows_of(rows, marker)
    }


def _judge_flagged(
    marker: str, wrong_values: list[str]
) -> Callable[[list[tuple]], tuple[bool, str]]:
    def judge(cells: list[tuple]) -> tuple[bool, str]:
        target = [row for row in cells if any(norm(marker) == norm(c) for c in row if c)]
        if not target:
            return False, f"{marker} の行が出力に無い"
        leftovers = [
            w for w in wrong_values if any(norm(w) == norm(c) for row in target for c in row if c)
        ]
        if leftovers:
            return False, f"確定した値として {'・'.join(leftovers)} が残っている"
        flagged = all(PENDING in " ".join(str(c) for c in row if c) for row in target)
        return flagged, f"{len(target)}行とも「{PENDING}」になっている" if flagged else "印が無い"

    return judge


CASES = {
    "c2": Case(
        "c2",
        "六角穴付きボルトの単位 m → 本",
        lambda rows: _fix_value(rows, "HBT-M8-25", "unit", {"value": "本"}),
        _unit_is_not("m"),
    ),
    "c3": Case(
        "c3",
        "型番 OSL-l520-N → OSL-1520-N",
        lambda rows: _fix_value(rows, "OSL-l520-N", "model_no", {"value": "OSL-1520-N"}),
        _absent("OSL-l520-N"),
    ),
    "c4": Case(
        "c4",
        "数量 2000（桁違い）→ 要確認",
        lambda rows: _fix_value(rows, "CPL-D40", "quantity", {"state": "needs_confirmation"}),
        _absent("2000"),
    ),
    "c5": Case(
        "c5",
        "型番の全角 → 半角（正規化後の値を人が確認）",
        lambda rows: _fix_value(rows, "ＨＢＴ－Ｍ１０－４０", "model_no", {"value": "HBT-M10-40"}),
        lambda cells: (
            "ＨＢＴ－Ｍ１０－４０" not in " ".join(str(c) for row in cells for c in row if c),
            "出力に全角の型番が無い",
        ),
    ),
    "c6": Case(
        "c6",
        "Oリングの単位 kg → 個",
        lambda rows: _fix_value(rows, "ORG-P18", "unit", {"value": "個"}),
        _unit_is_not("kg"),
    ),
    "c7": Case(
        "c7",
        "数量 40 / 400 の食い違い → 両方の行の数量を要確認",
        lambda rows: _flag_all(rows, "PPN-8-40", "quantity"),
        _judge_flagged("PPN-8-40", ["40", "400"]),
    ),
    "c8": Case(
        "c8",
        "納期 10/31 / 11/30 の食い違い → 両方の行の納期を要確認",
        lambda rows: _flag_all(rows, "ACY-SD40-150", "due_date"),
        _judge_flagged("ACY-SD40-150", ["2026-10-31", "2026-11-30"]),
    ),
}


def run(client: httpx.Client, case: Case, inquiry_id: str) -> tuple[bool, str, list[str]]:
    log: list[str] = []
    items = client.get(f"/api/v1/inquiries/{inquiry_id}/items").raise_for_status().json()
    log.append(
        f"読み取り結果: {len(items['rows'])}行 / 未確認 {items['summary']['unchecked_rows']}"
    )

    # 1. 誤りを直す（#10）。直した行はその場で確認済みになる
    for row_id, values in case.fix(items["rows"]).items():
        response = client.patch(
            f"/api/v1/inquiries/{inquiry_id}/items/{row_id}", json={"values": values}
        )
        response.raise_for_status()
        result = response.json()
        log.append(
            f"修正: {result['row_no']}行目 {list(values)} → {result['classification']}・確認済み"
        )

    # 2. 残りの行を確認する（#11）
    #    **確信が高い行は、読み取り元を開かないと確認できない**（② FUNC-04）。人と同じ手順で開く
    items = client.get(f"/api/v1/inquiries/{inquiry_id}/items").raise_for_status().json()
    remaining = [r for r in items["rows"] if r["check_state"] != "checked" and not r["excluded"]]
    opened = 0
    for row in remaining:
        if row["classification"] == "high_confidence" and not row["source_opened"]:
            value_id = next(
                (v["value_id"] for v in row["values"].values() if v.get("source")), None
            )
            if value_id:
                client.get(
                    f"/api/v1/inquiries/{inquiry_id}/values/{value_id}/source"
                ).raise_for_status()
                opened += 1
        client.post(
            f"/api/v1/inquiries/{inquiry_id}/items/{row['row_id']}/check"
        ).raise_for_status()
    log.append(f"確認: 読み取り元を {opened}行ぶん開き、残り {len(remaining)}行を確認済みにした")

    # 3. 確定する（#17）
    confirm = client.post(f"/api/v1/inquiries/{inquiry_id}/confirm")
    if confirm.status_code != 200:
        return False, f"確定できない: {confirm.json()['detail']}", log
    body = confirm.json()
    log.append(f"確定: {body['row_count']}行（顧客回答待ち {body['pending_row_count']}行）")

    # 4. 出力 Excel に誤りが残っていないか（#18）
    export = client.get(f"/api/v1/inquiries/{inquiry_id}/export").raise_for_status()
    path = Path(sys.argv[0]).parent / f"_kpi4_{case.key}.xlsx"
    path.write_bytes(export.content)
    sheet = load_workbook(path).active
    cells = [row for row in sheet.iter_rows(min_row=4, values_only=True)]
    path.unlink()
    ok, detail = case.judge(cells)
    return ok, detail, log


def main(targets: dict[str, str]) -> None:
    passed = 0
    with httpx.Client(base_url=BASE, timeout=60.0) as client:
        for key, inquiry_id in targets.items():
            case = CASES[key]
            print(f"\n=== {key}: {case.note}")
            try:
                ok, detail, log = run(client, case, inquiry_id)
            except Exception as e:  # noqa: BLE001 1件落ちても残りを続ける
                ok, detail, log = False, f"失敗: {e!r}", []
            for line in log:
                print(f"    {line}")
            print(f"    {'OK  ' if ok else 'NG  '}{detail}")
            passed += 1 if ok else 0
    print(f"\n見逃し0の検証: {passed} / {len(targets)} 件が合格")


if __name__ == "__main__":
    args = dict(a.split("=", 1) for a in sys.argv[1:] if "=" in a)
    if not args:
        raise SystemExit("使い方: uv run python scripts/verify_kpi4.py c2=<inquiry_id> ...")
    main(args)
